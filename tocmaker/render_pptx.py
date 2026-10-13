"""Scene -> editable PowerPoint (python-pptx).

Every panel background, label, badge and arrow becomes a native shape; icons are embedded as
SVG (PowerPoint 2016+ renders the vector and offers *Convert to Shape*) with a PNG fallback,
or optionally as EMF (converted with LibreOffice) which ungroups in any PowerPoint version.
The slide is exactly the TOC size (times ``scale``), so *File > Export* yields a correctly
sized image, and the whole picture can be copied into a bigger deck as one group.
"""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from typing import Any, Dict, List, Optional

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Pt

from . import icons as iconlib
from .layout import Scene
from .presets import Preset

EMU_PER_PT = 12700
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def _rgb(hexs: str) -> RGBColor:
    return RGBColor.from_string(hexs.lstrip("#").upper())


class PptxWriter:
    def __init__(self, preset: Preset, scale: float = 1.0, icon_format: str = "svg", font_name: str = "Arial",
                 notes: str = ""):
        self.preset, self.scale, self.icon_format, self.font_name = preset, scale, icon_format, font_name
        self.prs = Presentation()
        self.prs.slide_width = self._emu(preset.width_pt)
        self.prs.slide_height = self._emu(preset.height_pt)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self._n_media = 0
        self._tmp = tempfile.mkdtemp(prefix="tocmaker_pptx_")
        self.notes = notes
        self.warnings: List[str] = []

    def _emu(self, pt: float) -> Emu:
        return Emu(int(round(pt * self.scale * EMU_PER_PT)))

    # ------------------------------------------------------------------ primitives
    def write(self, scene: Scene, path: str) -> str:
        for e in scene:
            getattr(self, "_" + e["type"], self._unknown)(e)
        if self.notes:
            self.slide.notes_slide.notes_text_frame.text = self.notes
        self.prs.save(path)
        shutil.rmtree(self._tmp, ignore_errors=True)
        return path

    def _unknown(self, e):
        self.warnings.append(f"unknown primitive {e['type']}")

    def _rect(self, e):
        rx = e.get("rx", 0) or 0
        shp = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rx else MSO_SHAPE.RECTANGLE,
                                          self._emu(e["x"]), self._emu(e["y"]), self._emu(e["w"]), self._emu(e["h"]))
        if rx:
            shp.adjustments[0] = max(0.0, min(0.5, rx / max(1e-6, min(e["w"], e["h"]))))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(e["fill"])
        if e.get("stroke"):
            shp.line.color.rgb = _rgb(e["stroke"])
            shp.line.width = Pt(e.get("stroke_width", 0.5) * self.scale)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        shp.name = f"panel {e['panel']} ({e['role']})" if "panel" in e else ("background" if e["w"] >= self.preset.width_pt - 0.1 else "box")
        if e.get("placeholder"):
            shp.name = f"MISSING ICON: {e['placeholder']}"
            tf = shp.text_frame
            tf.text = e["placeholder"]
            tf.paragraphs[0].runs[0].font.size = Pt(4 * self.scale)

    def _text(self, e):
        n_lines = len(e["lines"])
        h = e["line_height"] * n_lines + 1.0
        tb = self.slide.shapes.add_textbox(self._emu(e["x"]), self._emu(e["y"] - 0.5), self._emu(e["w"]), self._emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Emu(0))
        tf.vertical_anchor = MSO_ANCHOR.TOP
        align = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[e.get("align", "center")]
        for i, ln in enumerate(e["lines"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = 1.0
            r = p.add_run()
            r.text = ln
            f = r.font
            f.size = Pt(e["size"] * self.scale)
            f.bold = e["weight"] == "bold"
            f.name = self.font_name
            f.color.rgb = _rgb(e["color"])
        tb.name = "label: " + e["lines"][0][:30]

    def _badge(self, e):
        shp = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, self._emu(e["x"]), self._emu(e["y"]),
                                          self._emu(e["w"]), self._emu(e["h"]))
        shp.adjustments[0] = 0.5
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(e["fill"])
        shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Emu(0))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        for i, ln in enumerate(e.get("lines") or [e["text"]]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 1.0
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(e["size"] * self.scale)
            r.font.bold = True
            r.font.name = self.font_name
            r.font.color.rgb = _rgb(e["color"])
        shp.name = "badge: " + e["text"][:30]

    def _line(self, e):
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self._emu(e["x1"]), self._emu(e["y1"]),
                                            self._emu(e["x2"]), self._emu(e["y2"]))
        c.line.color.rgb = _rgb(e["color"])
        c.line.width = Pt(e.get("width", 0.8) * self.scale)

    def _arrow(self, e):
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self._emu(e["x1"]), self._emu(e["y1"]),
                                            self._emu(e["x2"]), self._emu(e["y2"]))
        c.line.color.rgb = _rgb(e["color"])
        c.line.width = Pt(e.get("width", 1.4) * self.scale)
        if e.get("style") == "dashed":
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, f"{{{NS_A}}}tailEnd")
        if e.get("style") == "inhibit":
            tail.set("type", "none")
            # perpendicular bar at the head
            dx, dy = e["x2"] - e["x1"], e["y2"] - e["y1"]
            L = max(1e-6, (dx * dx + dy * dy) ** 0.5)
            nx, ny = -dy / L * 2.5, dx / L * 2.5
            bar = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self._emu(e["x2"] + nx), self._emu(e["y2"] + ny),
                                                  self._emu(e["x2"] - nx), self._emu(e["y2"] - ny))
            bar.line.color.rgb = _rgb(e["color"])
            bar.line.width = Pt(e.get("width", 1.4) * 1.3 * self.scale)
            bar.name = "inhibit bar"
        else:
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")
        c.name = "arrow" + (f": {e['label']}" if e.get("label") else "")
        if e.get("label"):
            mx, my = (e["x1"] + e["x2"]) / 2, (e["y1"] + e["y2"]) / 2
            horizontal = abs(e["x2"] - e["x1"]) >= abs(e["y2"] - e["y1"])
            w = max(20.0, abs(e["x2"] - e["x1"]) + 20)
            if horizontal:
                self._text(dict(x=mx - w / 2, y=my - e["size"] * 1.2 - 1.5, w=w, lines=[e["label"]], size=e["size"],
                                weight="regular", color=e["color"], align="center", line_height=e["size"] * 1.18))
            else:
                self._text(dict(x=mx + 2.5, y=my - e["size"] * 0.6, w=40, lines=[e["label"]], size=e["size"],
                                weight="regular", color=e["color"], align="left", line_height=e["size"] * 1.18))

    def _icon(self, e):
        info: iconlib.IconInfo = e["icon"]
        primary, accent, tint, flip = e["primary"], e["accent"], e.get("tint", "auto"), e.get("flip", False)
        png = iconlib.icon_png(info, 512, primary, accent, tint, flip)
        pic = self.slide.shapes.add_picture(io.BytesIO(png), self._emu(e["x"]), self._emu(e["y"]),
                                            self._emu(e["w"]), self._emu(e["h"]))
        pic.name = "icon: " + info.id
        if e.get("opacity", 1.0) < 1.0:
            self._set_alpha(pic, e["opacity"])
        self._n_media += 1
        svg_bytes = iconlib.icon_svg_bytes(info, primary, accent, tint, flip)
        if self.icon_format == "emf":
            emf = self._svg_to_emf(svg_bytes, self._n_media)
            if emf:
                rid = self._add_part(f"/ppt/media/tocicon{self._n_media}.emf", "image/x-emf", emf)
                blip = pic._element.blipFill.blip
                blip.set(f"{{{NS_R}}}embed", rid)
                return
            self.warnings.append("EMF conversion unavailable (LibreOffice not found); embedded SVG+PNG instead")
        rid = self._add_part(f"/ppt/media/tocicon{self._n_media}.svg", "image/svg+xml", svg_bytes)
        blip = pic._element.blipFill.blip
        ext_lst = etree.SubElement(blip, f"{{{NS_A}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{NS_A}}}ext")
        ext.set("uri", SVG_EXT_URI)
        sb = etree.SubElement(ext, f"{{{NS_ASVG}}}svgBlip")
        sb.set(f"{{{NS_R}}}embed", rid)

    def _set_alpha(self, pic, opacity: float):
        blip = pic._element.blipFill.blip
        a = etree.SubElement(blip, f"{{{NS_A}}}alphaModFix")
        a.set("amt", str(int(opacity * 100000)))

    def _add_part(self, partname: str, content_type: str, blob: bytes) -> str:
        part = Part(PackURI(partname), content_type, self.prs.part.package, blob)
        return self.slide.part.relate_to(part, RT.IMAGE)

    def _svg_to_emf(self, svg_bytes: bytes, n: int) -> Optional[bytes]:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            return None
        src = os.path.join(self._tmp, f"icon{n}.svg")
        with open(src, "wb") as fh:
            fh.write(svg_bytes)
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "emf", "--outdir", self._tmp, src],
                           capture_output=True, timeout=120, check=False)
        except Exception:
            return None
        out = os.path.join(self._tmp, f"icon{n}.emf")
        if os.path.exists(out):
            with open(out, "rb") as fh:
                return fh.read()
        return None


def render_pptx(scene: Scene, preset: Preset, path: str, scale: float = 1.0, icon_format: str = "svg",
                notes: str = "") -> List[str]:
    w = PptxWriter(preset, scale=scale, icon_format=icon_format, notes=notes)
    w.write(scene, path)
    return w.warnings
